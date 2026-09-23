import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

INPUT_STUDENTS_FILE = os.path.join("input", "students.xlsx")
BATCHES_FILE = os.path.join("input", "student_batches.xlsx")

def load_authoritative_students():
    """Reads Register Number and Student Name from input/students.xlsx (Account Details)."""
    students = []
    if not os.path.exists(INPUT_STUDENTS_FILE):
        return students

    wb = openpyxl.load_workbook(INPUT_STUDENTS_FILE, data_only=True)
    if "Account Details" not in wb.sheetnames:
        return students

    sheet = wb["Account Details"]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return students

    headers = [str(h).strip().upper() if h is not None else '' for h in rows[0]]
    reg_idx = -1
    name_idx = -1

    for idx, h in enumerate(headers):
        if "REGISTER" in h:
            reg_idx = idx
        elif "NAME" in h and name_idx == -1:
            name_idx = idx

    for row in rows[1:]:
        if not row or reg_idx >= len(row) or row[reg_idx] is None:
            continue
        reg_num = str(row[reg_idx]).strip()
        if not reg_num:
            continue
        name = str(row[name_idx]).strip() if name_idx != -1 and name_idx < len(row) and row[name_idx] is not None else ""
        students.append((reg_num, name))

    return students

def style_sheet(sheet, max_cols):
    """Applies header styling, freeze panes, autofilter, and column auto-width."""
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    sheet.freeze_panes = 'A2'

    # Header styling
    for col in range(1, max_cols + 1):
        cell = sheet.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align

    # Column width and borders
    for col in range(1, max_cols + 1):
        max_len = 0
        col_letter = get_column_letter(col)
        for row in range(1, sheet.max_row + 1):
            cell = sheet.cell(row=row, column=col)
            cell.border = thin_border
            if row > 1:
                # Text alignment
                cell.alignment = Alignment(vertical="center")
                # Format register numbers as text to prevent scientific notation
                if col == 1:
                    cell.number_format = '@'
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        sheet.column_dimensions[col_letter].width = max(max_len + 4, 15)

    if sheet.max_row >= 1 and max_cols >= 1:
        last_col_letter = get_column_letter(max_cols)
        sheet.auto_filter.ref = f"A1:{last_col_letter}{sheet.max_row}"

def create_or_refresh_batch_file():
    """
    Creates or updates input/student_batches.xlsx.
    Preserves existing Batch Details sheet if present.
    Updates Student Reference sheet with authoritative students.
    """
    students = load_authoritative_students()

    if os.path.exists(BATCHES_FILE):
        print(f"Loading existing batch file for reference update: {BATCHES_FILE}")
        wb = openpyxl.load_workbook(BATCHES_FILE)
    else:
        print(f"Creating new batch template file: {BATCHES_FILE}")
        wb = openpyxl.Workbook()
        # Remove default sheet
        wb.remove(wb.active)

    # 1. Ensure Batch Details sheet exists (preserve existing rows)
    if "Batch Details" not in wb.sheetnames:
        batch_sheet = wb.create_sheet(title="Batch Details", index=0)
        batch_sheet.append(["Register Number", "Student Name", "Batch"])
    else:
        batch_sheet = wb["Batch Details"]

    style_sheet(batch_sheet, 3)

    # 2. Re-create / Update Student Reference sheet
    if "Student Reference" in wb.sheetnames:
        wb.remove(wb["Student Reference"])

    ref_sheet = wb.create_sheet(title="Student Reference")
    ref_sheet.append(["Register Number", "Student Name"])

    for reg_num, name in students:
        ref_sheet.append([reg_num, name])

    style_sheet(ref_sheet, 2)

    os.makedirs(os.path.dirname(BATCHES_FILE), exist_ok=True)
    wb.save(BATCHES_FILE)
    print(f"Successfully saved {BATCHES_FILE} with {len(students)} student reference entries.")

if __name__ == "__main__":
    create_or_refresh_batch_file()
